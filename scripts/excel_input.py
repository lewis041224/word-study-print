"""Extract Excel cells without interpreting worksheet text as instructions."""
import re
from pathlib import Path


def extract_excel(path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    book = load_workbook(path, data_only=False, read_only=False)
    sheets = []
    try:
        for sheet in book.worksheets:
            rows = []
            for cells in sheet.iter_rows():
                if not any(c.value is not None for c in cells):
                    continue
                rows.append({'row': cells[0].row, 'cells': [
                    {'coordinate': c.coordinate, 'column': get_column_letter(c.column),
                     'value': str(c.value) if c.value is not None else '',
                     'formula': c.data_type == 'f'} for c in cells]})
            sheets.append({'sheet': sheet.title, 'rows': rows,
                           'merged_ranges': [str(r) for r in sheet.merged_cells.ranges]})
    finally:
        book.close()
    return {'source': Path(path).name, 'sheets': sheets,
            'needs_semantic_pairing': True, 'formulas_evaluated': False}


def import_pairs(path, pairs, sheet_name=None, start_row=2, numbered=False,
                 pos_column=None, unit_column=None, version='大学'):
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    pair_columns = []
    for pair in pairs.split(','):
        if not re.fullmatch(r'[A-Za-z]+:[A-Za-z]+', pair.strip()):
            raise ValueError('列对示例：A:B,C:D')
        pair_columns.append(tuple(column_index_from_string(c.upper()) for c in pair.strip().split(':')))
    if start_row < 1:
        raise ValueError('start-row 必须为正整数')
    book = load_workbook(path, read_only=False, data_only=False)
    sheets = [book[sheet_name]] if sheet_name else list(book.worksheets)
    entries = []
    try:
        for sheet in sheets:
            # Column pairs are separate source groups; keep each pair's order.
            for pair_index, (wc, mc) in enumerate(pair_columns):
                for row in range(start_row, sheet.max_row + 1):
                    wcell, mcell = sheet.cell(row, wc), sheet.cell(row, mc)
                    word = str(wcell.value).strip() if wcell.value is not None else ''
                    meaning = str(mcell.value).strip() if mcell.value is not None else ''
                    if not word and not meaning:
                        continue
                    if wcell.data_type == 'f' or mcell.data_type == 'f':
                        raise ValueError(f'{sheet.title} {wcell.coordinate}:{mcell.coordinate} 含公式；请先提供数值文本')
                    if not word or not meaning:
                        raise ValueError(f'{sheet.title} {wcell.coordinate}:{mcell.coordinate} 中英配对不完整')
                    if numbered:
                        wm = re.match(r'^\s*(\d+)[.、．]\s*(.+)$', word)
                        mm = re.match(r'^\s*(\d+)[.、．]\s*(.+)$', meaning)
                        if not wm or not mm or wm[1] != mm[1]:
                            raise ValueError(f'{sheet.title} 第{row}行两侧序号不一致')
                        word, meaning = wm[2], mm[2]
                    entry = {'word': word, 'meaning': meaning, 'book': Path(path).stem,
                             'unit': '', 'group_order': sheets.index(sheet) * len(pair_columns) + pair_index,
                             'order': len(entries), 'reviewed': False,
                             'source': f'{Path(path).name} {sheet.title} {wcell.coordinate}:{mcell.coordinate}'}
                    for key, column in [('pos', pos_column), ('unit', unit_column)]:
                        if column:
                            cell = sheet.cell(row, column_index_from_string(column.upper()))
                            if cell.data_type == 'f':
                                raise ValueError(f'{cell.coordinate} 含公式')
                            entry[key] = '' if cell.value is None else str(cell.value).strip()
                    entries.append(entry)
    finally:
        book.close()
    if not entries:
        raise ValueError('指定范围未找到词汇')
    return {'settings': {'version': version, 'preserve_source_order': True}, 'entries': entries,
            'needs_semantic_review': True}
