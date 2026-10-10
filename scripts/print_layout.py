"""Day-first printable vocabulary tables with natural Word pagination."""
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def new_document(header='背单词练习'):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.7)
    s.left_margin = s.right_margin = Cm(1.8)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(16 if name == 'Normal' else 18)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '仿宋' if name == 'Normal' else '黑体')
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.05
    p = s.header.paragraphs[0]
    p.text = header
    for run in p.runs:
        run.font.size = Pt(12)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '仿宋')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    s.footer.paragraphs[0]._p.append(field)
    return doc


def paragraph(doc, text, style='Heading 1', new_page=False, small=False):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.page_break_before = new_page
    if small:
        for run in p.runs:
            run.font.size = Pt(12)
    return p


def table(doc, rows, feedback=False, memory=False):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column in t.columns:
        column.width = Cm(8.7)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
            border.set(qn('w:' + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for index, pair in enumerate(rows):
        row = t.add_row()
        row.height = Cm(.7 if index == 0 else (1.3 if feedback else (.82 if memory else 1.0)))
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for cell, text in zip(row.cells, pair):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F2F2F2')
                cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                # A row may wrap but must not force the following row to stay with it.
                p.paragraph_format.keep_with_next = False
                if index == 0:
                    for run in p.runs:
                        run.font.size = Pt(12)
    return t


def chinese(entry):
    pos = str(entry.get('pos', '')).strip()
    return (pos + ' ' if pos else '') + entry['meaning']


def day_summary(day, learned_on):
    label = day.get('date') or f"第{day['day']}天"
    review_days = sorted({learned_on[i] for i in day['review']})
    review = ('第' + '、'.join(map(str, review_days)) + '天词汇') if review_days else '无'
    surprise = f"；穿插已学词{len(day.get('surprise', []))}个" if day.get('surprise') else ''
    return f"{label}：新学{len(day['new'])}词；复习{review}，共{len(day['review'])}词{surprise}。"


def practice_document(plan, days, target, header='背单词练习'):
    doc = new_document(header)
    entries = plan['entries']
    learned_on = {i: day['day'] for day in plan['days'] for i in day['new']}
    started = False
    for day in days:
        ids = day['entries']
        for mode, title, left, right in ((0, '中英对照', '英文', '中文'),
                                        (1, '中译英', '中文', '英文'),
                                        (2, '英译中', '英文', '中文')):
            paragraph(doc, f"第{day['day']}天 {title}", new_page=started)
            paragraph(doc, day_summary(day, learned_on), 'Normal', small=True)
            for offset in range(0, len(ids), 20):
                if offset:
                    paragraph(doc, f"第{day['day']}天 {title} 续", new_page=True)
                rows = [(left + f" · 第{day['day']}天", right + ' · ' + title)]
                for i in ids[offset:offset+20]:
                    entry = entries[i]
                    en, zh = entry['word'], chinese(entry)
                    pos = entry.get('pos', '') if plan['settings'].get('translation_pos_hint', True) else ''
                    rows.append((en, zh) if mode == 0 else ((zh, '') if mode == 1 else (en, pos)))
                table(doc, rows)
            started = True
        paragraph(doc, f"第{day['day']}天 反馈表", new_page=True)
        table(doc, [('英文', '中文')] + [('', '') for _ in range(14)], feedback=True)
    doc.save(target)


def memory_document(plan, target, title='单词总版记忆表'):
    doc = new_document(title)
    paragraph(doc, title, 'Title')
    table(doc, [('英文', '中文')] + [(e['word'], chinese(e)) for e in plan['entries']], memory=True)
    doc.save(target)
